#!/usr/bin/env python3
"""
kca_acs.py - Kansas City ACS 5-year comparison builder (2014-2024).

Pulls American Community Survey 5-year estimates from the Census Data API for
four Kansas City geographies, computes derived metrics, and writes:

  kc_acs_2014_2024.csv    long format: geography, year, variable, value
  kc_acs_2014_2024.xlsx   one sheet per geography + Summary + Notes
  charts/*.png            one line chart per metric, all geographies
  kc_acs_summary.md       plain-English trend bullets for an acquisitions team

Usage
  export CENSUS_API_KEY=...            # never hardcode the key
  python kca_acs.py --test             # one call (2024, KCMO), print raw values
  python kca_acs.py                    # full 2014-2024 build
  python kca_acs.py --out-dir out/     # write outputs elsewhere

Requires: requests, pandas, openpyxl, matplotlib
"""
from __future__ import annotations

import argparse
import logging
import os
import sys
import time
from pathlib import Path

import pandas as pd
import requests

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #
API_URL = "https://api.census.gov/data/{year}/acs/acs5"
YEARS = list(range(2014, 2025))  # 11 five-year releases: 2010-14 ... 2020-24
BASE_YEAR, END_YEAR = YEARS[0], YEARS[-1]
TIMEOUT = 30
RETRY_PAUSE = 2  # seconds before the single retry

# label -> (for clause, in clause or None)
GEOGRAPHIES: dict[str, tuple[str, str | None]] = {
    "Kansas City city, MO": ("place:38000", "state:29"),
    "Kansas City MO-KS Metro": (
        "metropolitan statistical area/micropolitan statistical area:28140",
        None,
    ),
    "Johnson County, KS": ("county:091", "state:20"),
    "Jackson County, MO": ("county:095", "state:29"),
}

VARIABLES = {
    "B01003_001E": "total_population",
    "B25003_001E": "occupied_units",
    "B25003_003E": "renter_occupied_units",
    "B19013_001E": "median_hh_income",
    "B25077_001E": "median_home_value",
    "B03002_001E": "race_universe_pop",
    "B03002_003E": "white_nh_pop",
}

# CPI-U, U.S. city average, all items, 1982-84=100, not seasonally adjusted,
# annual averages (BLS series CUUR0000SA0).
CPI_U = {
    2014: 236.736,
    2015: 237.017,
    2016: 240.007,
    2017: 245.120,
    2018: 251.107,
    2019: 255.657,
    2020: 258.811,
    2021: 270.970,
    2022: 292.655,
    2023: 304.702,
    2024: 313.689,
}

# Wide-table column order and display formats (used by the xlsx writer).
METRICS = [
    # column,                    header,                                 format
    ("total_population",         "Total population",                     "#,##0"),
    ("occupied_units",           "Occupied housing units",               "#,##0"),
    ("renter_occupied_units",    "Renter-occupied units",                "#,##0"),
    ("renter_share",             "Renter share",                         "0.0%"),
    ("median_hh_income",         "Median HH income (nominal $)",         "$#,##0"),
    ("median_hh_income_2024usd", f"Median HH income ({END_YEAR} $)",     "$#,##0"),
    ("median_home_value",        "Median home value (nominal $)",        "$#,##0"),
    ("median_home_value_2024usd", f"Median home value ({END_YEAR} $)",   "$#,##0"),
    ("race_universe_pop",        "Population (race/ethnicity universe)", "#,##0"),
    ("white_nh_pop",             "White alone, not Hispanic",            "#,##0"),
    ("white_nh_share",           "White non-Hispanic share",             "0.0%"),
]
METRIC_HEADERS = {col: hdr for col, hdr, _ in METRICS}
METRIC_FORMATS = {col: fmt for col, _, fmt in METRICS}

# Metrics that get 2014 -> 2024 change in the Summary sheet.
# Shares get percentage-point change instead of % change.
SUMMARY_METRICS = [
    "total_population",
    "median_hh_income",
    "median_hh_income_2024usd",
    "median_home_value",
    "median_home_value_2024usd",
    "renter_share",
    "white_nh_share",
]
SHARE_METRICS = {"renter_share", "white_nh_share"}

# One PNG per metric.
CHART_METRICS = [
    "total_population",
    "total_population_index",
    "renter_share",
    "white_nh_share",
    "median_hh_income",
    "median_hh_income_2024usd",
    "median_home_value",
    "median_home_value_2024usd",
]

METRO_NOTE = (
    "The Kansas City, MO-KS metro (CBSA 28140) is reported on the OMB "
    "delineation in effect for each ACS release. OMB updated CBSA delineations "
    "several times during 2014-2024 (2013, 2015, 2017, 2018, 2020 and 2023 "
    "bulletins), so the set of counties in the metro is not fixed. Metro totals "
    "are therefore not perfectly comparable year to year."
)
OVERLAP_NOTE = (
    "ACS 5-year estimates are period estimates (e.g. 2024 = 2020-2024). "
    "Consecutive releases share four of five sample years, so year-over-year "
    "changes are smoothed; 2014 (2010-14) and 2024 (2020-24) share no years."
)
CPI_NOTE = (
    "Inflation adjustment: value x CPI-U(2024) / CPI-U(release year), using "
    "BLS CPI-U annual averages (U.S. city average, all items, 1982-84=100). "
    "ACS 5-year dollar figures are published in dollars of the release's final "
    "year, so the release year is the right base. Note the Census Bureau itself "
    "uses CPI-U-RS for its own adjustments, so results differ slightly."
)

log = logging.getLogger("kca_acs")


# --------------------------------------------------------------------------- #
# Fetch
# --------------------------------------------------------------------------- #
def _api_key() -> str | None:
    key = os.environ.get("CENSUS_API_KEY", "").strip()
    if not key:
        log.warning(
            "CENSUS_API_KEY is not set; calling the API without a key "
            "(limited to 500 queries/day per IP)."
        )
        return None
    return key


def fetch_one(year: int, geo_label: str, key: str | None,
              session: requests.Session) -> dict | None:
    """Fetch all VARIABLES for one geography/year. Retries once; returns a
    dict of raw strings keyed by variable code, or None on failure."""
    geo_for, geo_in = GEOGRAPHIES[geo_label]
    params = {"get": ",".join(["NAME", *VARIABLES]), "for": geo_for}
    if geo_in:
        params["in"] = geo_in
    if key:
        params["key"] = key

    url = API_URL.format(year=year)
    last_err = None
    for attempt in (1, 2):
        try:
            resp = session.get(url, params=params, timeout=TIMEOUT)
            resp.raise_for_status()
            # Census returns HTML/plain text (status 200) for some errors,
            # e.g. an invalid key, so a JSON decode failure counts as failure.
            rows = resp.json()
            if len(rows) < 2:
                raise ValueError("empty result set")
            return dict(zip(rows[0], rows[1]))
        except (requests.RequestException, ValueError) as err:
            last_err = err
            if attempt == 1:
                time.sleep(RETRY_PAUSE)
    log.error("FAILED %s %s after retry: %s", geo_label, year, last_err)
    return None


def to_number(raw) -> float | None:
    """Convert a Census value to float. Null and the negative annotation
    sentinels (-666666666, -999999999, -888888888, -222222222, ...) become
    None. All variables in this script are counts or medians, so any
    negative value is a sentinel."""
    if raw is None:
        return None
    try:
        val = float(raw)
    except (TypeError, ValueError):
        return None
    return None if val < 0 else val


def fetch_all(years=YEARS, geographies=GEOGRAPHIES):
    """Returns (records, failures). records: list of dicts with geography,
    year, census_name and one key per friendly variable name."""
    key = _api_key()
    records, failures = [], []
    with requests.Session() as session:
        for year in years:
            for geo in geographies:
                raw = fetch_one(year, geo, key, session)
                if raw is None:
                    failures.append((geo, year))
                    continue
                rec = {"geography": geo, "year": year,
                       "census_name": raw.get("NAME")}
                for code, name in VARIABLES.items():
                    rec[name] = to_number(raw.get(code))
                records.append(rec)
                log.info("ok  %s %s (%s)", geo, year, rec["census_name"])
    return records, failures


# --------------------------------------------------------------------------- #
# Transform
# --------------------------------------------------------------------------- #
def _safe_div(num: pd.Series, den: pd.Series) -> pd.Series:
    return num.where(den > 0) / den.where(den > 0)


def build_wide(records: list[dict]) -> pd.DataFrame:
    """Geography x year table with raw and derived metrics."""
    cols = ["geography", "year", "census_name", *VARIABLES.values()]
    df = pd.DataFrame.from_records(records, columns=cols)
    df["renter_share"] = _safe_div(df["renter_occupied_units"], df["occupied_units"])
    df["white_nh_share"] = _safe_div(df["white_nh_pop"], df["race_universe_pop"])
    factor = df["year"].map(lambda y: CPI_U[END_YEAR] / CPI_U[y])
    df["median_hh_income_2024usd"] = df["median_hh_income"] * factor
    df["median_home_value_2024usd"] = df["median_home_value"] * factor
    df["geography"] = pd.Categorical(df["geography"], categories=list(GEOGRAPHIES))
    return df.sort_values(["geography", "year"]).reset_index(drop=True)


def build_long(wide: pd.DataFrame) -> pd.DataFrame:
    metric_cols = [c for c, _, _ in METRICS]
    long = wide.melt(id_vars=["geography", "year"], value_vars=metric_cols,
                     var_name="variable", value_name="value")
    order = {c: i for i, c in enumerate(metric_cols)}
    long["_o"] = long["variable"].map(order)
    long = long.sort_values(["geography", "year", "_o"]).drop(columns="_o")
    long["geography"] = long["geography"].astype(str)
    return long.reset_index(drop=True)


def build_summary(wide: pd.DataFrame) -> pd.DataFrame:
    """2014 vs 2024 per geography per metric, with absolute and % change
    (percentage-point change for shares)."""
    rows = []
    for geo in GEOGRAPHIES:
        g = wide[wide["geography"] == geo].set_index("year")
        for m in SUMMARY_METRICS:
            v0 = g[m].get(BASE_YEAR) if BASE_YEAR in g.index else None
            v1 = g[m].get(END_YEAR) if END_YEAR in g.index else None
            v0 = None if pd.isna(v0) else v0
            v1 = None if pd.isna(v1) else v1
            change = (v1 - v0) if v0 is not None and v1 is not None else None
            pct = (change / v0) if change is not None and v0 and m not in SHARE_METRICS else None
            rows.append({
                "geography": geo,
                "metric": METRIC_HEADERS[m],
                "metric_key": m,
                str(BASE_YEAR): v0,
                str(END_YEAR): v1,
                "abs_change": change,
                "pct_change": pct,
            })
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def write_excel(wide, summary, failures, path: Path) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")
    note_font = Font(italic=True, color="555555", size=9)

    def style_header(ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center",
                                       horizontal="center")
        ws.row_dimensions[row].height = 45

    def sheet_name(geo):  # Excel: <=31 chars, no : \ / ? * [ ]
        return "".join(ch for ch in geo if ch not in ':\\/?*[]').replace(",", "")[:31]

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        # ---- Summary ------------------------------------------------------
        summ = summary.drop(columns="metric_key").rename(columns={
            "geography": "Geography", "metric": "Metric",
            "abs_change": f"Change {BASE_YEAR}->{END_YEAR}",
            "pct_change": "% change",
        })
        summ.to_excel(xw, sheet_name="Summary", index=False, startrow=0)
        ws = xw.sheets["Summary"]
        style_header(ws, 1, summ.shape[1])
        for i, key in enumerate(summary["metric_key"], start=2):
            fmt = METRIC_FORMATS[key]
            for col in (3, 4):
                ws.cell(row=i, column=col).number_format = fmt
            # shares: change is percentage points, shown as e.g. +3.1 pp
            ws.cell(row=i, column=5).number_format = (
                '+0.0" pp";-0.0" pp"' if key in SHARE_METRICS else
                "+$#,##0;-$#,##0" if "$" in fmt else "+#,##0;-#,##0"
            )
            if key in SHARE_METRICS:
                c = ws.cell(row=i, column=5)
                c.value = None if c.value is None else c.value * 100
            ws.cell(row=i, column=6).number_format = "+0.0%;-0.0%"
        for col, width in zip("ABCDEF", (26, 34, 14, 14, 18, 11)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "C2"
        r = len(summ) + 3
        for note in ("Notes:", METRO_NOTE, OVERLAP_NOTE, CPI_NOTE,
                     "Share changes are in percentage points (pp)."):
            ws.cell(row=r, column=1, value=note).font = note_font
            r += 1

        # ---- One sheet per geography -------------------------------------
        cols = [c for c, _, _ in METRICS]
        for geo in GEOGRAPHIES:
            g = wide[wide["geography"] == geo][["year", *cols]].copy()
            g = g.rename(columns={"year": "Year", **METRIC_HEADERS})
            name = sheet_name(geo)
            g.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            style_header(ws, 1, g.shape[1])
            for j, col in enumerate(cols, start=2):
                for i in range(2, len(g) + 2):
                    ws.cell(row=i, column=j).number_format = METRIC_FORMATS[col]
                ws.column_dimensions[get_column_letter(j)].width = 16
            ws.column_dimensions["A"].width = 8
            ws.freeze_panes = "B2"
            r = len(g) + 3
            names = wide.loc[wide["geography"] == geo, "census_name"].dropna().unique()
            notes = [f"Census name(s): {'; '.join(names)}", OVERLAP_NOTE, CPI_NOTE]
            if "Metro" in geo:
                notes.insert(0, "CAUTION: " + METRO_NOTE)
            missing = [y for gg, y in failures if gg == geo]
            if missing:
                notes.insert(0, f"Failed requests (no data): {missing}")
            for note in notes:
                ws.cell(row=r, column=1, value=note).font = note_font
                r += 1

        # ---- Notes / CPI table --------------------------------------------
        cpi = pd.DataFrame({
            "Year": list(CPI_U),
            "CPI-U annual avg": list(CPI_U.values()),
            f"Factor to {END_YEAR} $": [CPI_U[END_YEAR] / v for v in CPI_U.values()],
        })
        cpi.to_excel(xw, sheet_name="Notes", index=False)
        ws = xw.sheets["Notes"]
        style_header(ws, 1, 3)
        for i in range(2, len(cpi) + 2):
            ws.cell(row=i, column=2).number_format = "0.000"
            ws.cell(row=i, column=3).number_format = "0.0000"
        for col, w in zip("ABC", (8, 16, 16)):
            ws.column_dimensions[col].width = w
        r = len(cpi) + 3
        lines = [
            "Source: U.S. Census Bureau, ACS 5-year estimates, "
            "https://api.census.gov/data/{YEAR}/acs/acs5",
            "Variables: " + ", ".join(f"{k} ({v})" for k, v in VARIABLES.items()),
            "CPI source: BLS CPI-U, series CUUR0000SA0, annual average.",
            METRO_NOTE, OVERLAP_NOTE, CPI_NOTE,
            "Missing values: Census null / negative sentinel codes "
            "(e.g. -666666666) are left blank.",
            "Failed requests: " + (", ".join(f"{g} {y}" for g, y in failures)
                                   if failures else "none"),
        ]
        for line in lines:
            ws.cell(row=r, column=1, value=line).font = note_font
            r += 1


# --------------------------------------------------------------------------- #
# Charts
# --------------------------------------------------------------------------- #
# Categorical slots 1-4 of the validated default palette, fixed per geography.
SERIES_COLORS = dict(zip(GEOGRAPHIES, ["#2a78d6", "#eb6834", "#1baf7a", "#eda100"]))
SURFACE, INK, INK_2, GRID = "#fcfcfb", "#0b0b0b", "#52514e", "#e4e3df"
CHART_TITLES = {
    "total_population": "Total population",
    "total_population_index": f"Population growth, indexed ({BASE_YEAR} = 100)",
    "renter_share": "Renter share of occupied housing units",
    "white_nh_share": "White non-Hispanic share of population",
    "median_hh_income": "Median household income (nominal dollars)",
    "median_hh_income_2024usd": f"Median household income ({END_YEAR} dollars, CPI-U)",
    "median_home_value": "Median home value, owner-occupied (nominal dollars)",
    "median_home_value_2024usd": f"Median home value, owner-occupied ({END_YEAR} dollars, CPI-U)",
}


def write_charts(wide: pd.DataFrame, out_dir: Path) -> list[Path]:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter, PercentFormatter

    out_dir.mkdir(parents=True, exist_ok=True)
    wide = wide.copy()
    base = wide[wide["year"] == BASE_YEAR].set_index("geography")["total_population"]
    wide["total_population_index"] = (
        wide["total_population"] / wide["geography"].map(base).astype(float) * 100)
    paths = []
    for m in CHART_METRICS:
        fig, ax = plt.subplots(figsize=(10, 5.6), dpi=150)
        fig.patch.set_facecolor(SURFACE)
        ax.set_facecolor(SURFACE)
        ends = []
        for geo in GEOGRAPHIES:
            # Reindex to every year so missing data shows as a gap, not a bridge.
            g = (wide[wide["geography"] == geo].set_index("year")[m]
                 .reindex(YEARS).astype(float))
            if g.isna().all():
                continue
            color = SERIES_COLORS[geo]
            ax.plot(g.index, g.values, color=color, linewidth=2, label=geo,
                    marker="o", markersize=4, solid_capstyle="round")
            last = g.dropna()
            ends.append((last.iloc[-1], geo, color, last.index[-1]))

        if m == "total_population_index":
            fmt = lambda v: f"{v:.1f}"
        elif m in SHARE_METRICS:
            ax.yaxis.set_major_formatter(PercentFormatter(1.0, decimals=0))
            fmt = lambda v: f"{v:.1%}"
        elif "income" in m or "value" in m:
            ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v/1000:,.0f}K"))
            fmt = lambda v: f"${v/1000:,.0f}K"
        else:
            ax.yaxis.set_major_formatter(FuncFormatter(
                lambda v, _: f"{v/1e6:.1f}M" if v >= 1e6 else f"{v/1000:,.0f}K"))
            fmt = lambda v: f"{v/1e6:.2f}M" if v >= 1e6 else f"{v/1000:,.0f}K"

        # Direct end labels in ink color (colored dot carries identity),
        # nudged apart so they never collide.
        ends.sort(key=lambda t: t[0])
        lo, hi = ax.get_ylim()
        min_gap = (hi - lo) * 0.055
        placed = []
        for val, *_ in ends:
            placed.append(val if not placed or val - placed[-1] >= min_gap
                          else placed[-1] + min_gap)
        if ends:  # re-center the stack on the data, keep it inside the axes
            shift = (sum(placed) - sum(e[0] for e in ends)) / len(ends)
            shift = max(shift, placed[-1] - hi + min_gap / 2)
            shift = min(shift, placed[0] - lo - min_gap / 2)
            placed = [y - shift for y in placed]
        for (val, geo, color, yr), y in zip(ends, placed):
            ax.annotate(f"{geo}  {fmt(val)}", xy=(yr, val), xytext=(yr + 0.25, y),
                        textcoords="data", va="center", fontsize=8.5, color=INK,
                        arrowprops=dict(arrowstyle="-", color=color, lw=0.8,
                                            relpos=(0, 0.5), shrinkA=2))

        ax.set_xlim(YEARS[0] - 0.4, YEARS[-1] + 3.6)
        ax.set_xticks(YEARS)
        ax.set_xlabel("ACS 5-year release (final year)", color=INK_2, fontsize=9)
        ax.set_title(CHART_TITLES[m], loc="left", fontsize=13, color=INK, pad=30)
        ax.grid(axis="y", color=GRID, linewidth=0.8)
        ax.tick_params(colors=INK_2, labelsize=9, length=0)
        for side in ("top", "right", "left"):
            ax.spines[side].set_visible(False)
        ax.spines["bottom"].set_color(GRID)
        leg = ax.legend(loc="lower left", bbox_to_anchor=(0, 1.0), fontsize=8.5,
                        frameon=False, ncol=4, handlelength=1.6, borderaxespad=0.3)
        for t in leg.get_texts():
            t.set_color(INK)
        foot = "Source: U.S. Census Bureau, ACS 5-year estimates."
        if any("Metro" in e[1] for e in ends):
            foot += " Metro boundary follows the OMB delineation in effect each year; not perfectly comparable."
        if m.endswith("2024usd"):
            foot += f" Adjusted with CPI-U annual averages to {END_YEAR} dollars."
        fig.text(0.01, 0.01, foot, fontsize=7.5, color=INK_2)
        fig.tight_layout(rect=(0, 0.03, 1, 1))
        path = out_dir / f"{m}.png"
        fig.savefig(path, facecolor=SURFACE)
        plt.close(fig)
        paths.append(path)
    return paths


# --------------------------------------------------------------------------- #
# Plain-English summary
# --------------------------------------------------------------------------- #
def write_summary_md(summary: pd.DataFrame, failures, path: Path) -> str:
    """Data-driven draft of the 5 trend bullets. Review before circulating."""
    s = summary.set_index(["geography", "metric_key"])

    def v(geo, m, col):
        try:
            x = s.loc[(geo, m), col]
        except KeyError:
            return None
        return None if pd.isna(x) else x

    def pct(x):
        return "n/a" if x is None else f"{x:+.1%}"

    def share(x):
        return "n/a" if x is None else f"{x:.1%}"

    def pp(x):
        return "n/a" if x is None else f"{x*100:+.1f} pp"

    b0, b1 = str(BASE_YEAR), str(END_YEAR)

    renter = "; ".join(
        f"{g} {share(v(g, 'renter_share', b0))} -> {share(v(g, 'renter_share', b1))} "
        f"({pp(v(g, 'renter_share', 'abs_change'))})" for g in GEOGRAPHIES)
    pop = "; ".join(f"{g} {pct(v(g, 'total_population', 'pct_change'))}"
                    for g in GEOGRAPHIES)
    inc_vs_home = "; ".join(
        f"{g}: income {pct(v(g, 'median_hh_income', 'pct_change'))} vs "
        f"home value {pct(v(g, 'median_home_value', 'pct_change'))}"
        for g in GEOGRAPHIES)
    real = "; ".join(
        f"{g}: real income {pct(v(g, 'median_hh_income_2024usd', 'pct_change'))}, "
        f"real home value {pct(v(g, 'median_home_value_2024usd', 'pct_change'))}"
        for g in GEOGRAPHIES)
    demo = "; ".join(
        f"{g} {share(v(g, 'white_nh_share', b0))} -> {share(v(g, 'white_nh_share', b1))} "
        f"({pp(v(g, 'white_nh_share', 'abs_change'))})" for g in GEOGRAPHIES)

    bullets = [
        f"**Renter share ({b0} -> {b1}):** {renter}.",
        f"**Population growth ({b0} -> {b1}):** {pop}.",
        f"**Income vs. home value, nominal:** {inc_vs_home}.",
        f"**Same comparison in {b1} dollars:** {real}.",
        f"**Demographic shift, White non-Hispanic share:** {demo}.",
    ]
    text = (
        f"# Kansas City ACS 5-year trends, {b0}-{b1}: data summary\n\n"
        "_Auto-generated from the pulled data. Rewrite into prose before "
        "circulating._\n\n"
        + "\n".join(f"- {b}" for b in bullets)
        + f"\n\n> {METRO_NOTE}\n>\n> {OVERLAP_NOTE}\n"
        + (f"\n**Failed requests:** {', '.join(f'{g} {y}' for g, y in failures)}\n"
           if failures else "")
    )
    path.write_text(text, encoding="utf-8")
    return text


# --------------------------------------------------------------------------- #
# Entry points
# --------------------------------------------------------------------------- #
def run_test(year: int = END_YEAR, geo: str = "Kansas City city, MO") -> int:
    """Single sanity-check call; prints raw API values."""
    key = _api_key()
    with requests.Session() as session:
        raw = fetch_one(year, geo, key, session)
    if raw is None:
        return 1
    print(f"\nRaw ACS 5-year values, {year}, {raw.get('NAME')}")
    print("-" * 64)
    for code, name in VARIABLES.items():
        r = raw.get(code)
        n = to_number(r)
        shown = "MISSING" if n is None else f"{n:,.0f}"
        print(f"{code:<13} {name:<24} raw={r!s:<12} -> {shown}")
    occ, rent = to_number(raw.get("B25003_001E")), to_number(raw.get("B25003_003E"))
    tot, wnh = to_number(raw.get("B03002_001E")), to_number(raw.get("B03002_003E"))
    if occ and rent is not None:
        print(f"{'':13} {'renter_share':<24} {rent / occ:.1%}")
    if tot and wnh is not None:
        print(f"{'':13} {'white_nh_share':<24} {wnh / tot:.1%}")
    return 0


def run_full(out_dir: Path) -> int:
    out_dir.mkdir(parents=True, exist_ok=True)
    records, failures = fetch_all()
    if not records:
        log.error("No data retrieved; nothing to write.")
        return 1
    wide = build_wide(records)
    long = build_long(wide)
    summary = build_summary(wide)

    csv_path = out_dir / "kc_acs_2014_2024.csv"
    xlsx_path = out_dir / "kc_acs_2014_2024.xlsx"
    md_path = out_dir / "kc_acs_summary.md"
    long.to_csv(csv_path, index=False)
    write_excel(wide, summary, failures, xlsx_path)
    charts = write_charts(wide, out_dir / "charts")
    write_summary_md(summary, failures, md_path)

    log.info("wrote %s (%d rows)", csv_path, len(long))
    log.info("wrote %s", xlsx_path)
    log.info("wrote %d charts to %s", len(charts), out_dir / "charts")
    log.info("wrote %s", md_path)
    if failures:
        log.warning("%d geography/year requests failed: %s", len(failures), failures)
    log.warning("NOTE: %s", METRO_NOTE)
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--test", action="store_true",
                    help="single sanity-check call (2024, Kansas City city, MO)")
    ap.add_argument("--out-dir", type=Path, default=Path("."),
                    help="output directory (default: current directory)")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    return run_test() if args.test else run_full(args.out_dir)


if __name__ == "__main__":
    sys.exit(main())
